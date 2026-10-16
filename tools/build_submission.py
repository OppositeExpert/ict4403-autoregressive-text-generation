"""Build the assignment Word report from actual experiment artifacts.

Usage: python tools/build_submission.py --name NAME --registration REGNUM
Requires python-docx; execute from the project root. Export the resulting DOCX
to PDF with Word or LibreOffice. Personal fields are never hardcoded here.
"""
import argparse
import csv
import json
import re
import shutil
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

LABELS={"rnn":"Simple RNN","birnn":"Bidirectional RNN","lstm":"LSTM","gru":"GRU"}


def csv_rows(path):
    with Path(path).open(newline="") as f:
        return list(csv.DictReader(f))


def link(paragraph, label, url):
    rel=paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement("w:hyperlink");h.set(qn("r:id"),rel)
    r=OxmlElement("w:r");properties=OxmlElement("w:rPr")
    color=OxmlElement("w:color");color.set(qn("w:val"),"245B87");properties.append(color)
    r.append(properties);t=OxmlElement("w:t");t.text=label;r.append(t);h.append(r)
    paragraph._p.append(h)


def table(document, headers, rows, widths):
    t=document.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for cell,width in zip(t.columns,widths):
        cell.width=Inches(width)
    for i,header in enumerate(headers):
        t.rows[0].cells[i].text=header
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=str(value)
    for rownum,row in enumerate(t.rows):
        props=row._tr.get_or_add_trPr()
        cant=OxmlElement("w:cantSplit");props.append(cant)
        if rownum==0:
            repeat=OxmlElement("w:tblHeader");props.append(repeat)
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right"]:
                e=OxmlElement("w:"+edge);e.set(qn("w:val"),"single");e.set(qn("w:sz"),"4");e.set(qn("w:color"),"D9D9D9");borders.append(e)
            pr.append(borders)
            margin=OxmlElement("w:tcMar")
            for edge,value in [("top",95),("bottom",95),("left",100),("right",100)]:
                e=OxmlElement("w:"+edge);e.set(qn("w:w"),str(value));e.set(qn("w:type"),"dxa");margin.append(e)
            pr.append(margin)
            fill=OxmlElement("w:shd");fill.set(qn("w:fill"),"244A66" if rownum==0 else ("F1F5F8" if rownum%2 else "FFFFFF"));pr.append(fill)
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
                for run in p.runs:
                    run.font.size=Pt(9.4);run.font.bold=(rownum==0)
                    run.font.color.rgb=RGBColor.from_string("FFFFFF" if rownum==0 else "000000")
    document.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


def paragraph(d,text,style=None):
    return d.add_paragraph(text,style)


def heading(d,text):
    d.add_heading(text,level=1)


def figure(d,path,alt,width=6.65):
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(3)
    shape=p.add_run().add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set("descr",alt)


def caption(d,text):
    p=d.add_paragraph(text)
    p.paragraph_format.space_after=Pt(8)
    for r in p.runs:
        r.font.size=Pt(9);r.italic=True


def main():
    p=argparse.ArgumentParser()
    p.add_argument("--name",required=True)
    p.add_argument("--registration",required=True)
    p.add_argument("--repository",default="https://github.com/OppositeExpert/ict4403-autoregressive-text-generation")
    p.add_argument("--results",default="results")
    p.add_argument("--output",default="submission")
    a=p.parse_args();root=Path(a.results);out=Path(a.output);out.mkdir(parents=True,exist_ok=True)
    summaries=csv_rows(root/"summary.csv")
    meta=json.loads((root/"dataset.json").read_text())
    manifest=json.loads((root/"run_manifest.json").read_text())
    verification=json.loads((root/"verification.json").read_text())
    samples=json.loads((root/"samples.json").read_text())
    baseline=json.loads((root/"baseline.json").read_text())
    temperatures=csv_rows(root/"temperature_summary.csv")
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.27);s.page_height=Inches(11.69)
    s.top_margin=s.bottom_margin=Inches(.68);s.left_margin=s.right_margin=Inches(.8)
    s.header_distance=s.footer_distance=Inches(.3)
    for name in ["Normal","Title","Subtitle","Heading 1","Heading 2"]:
        style=d.styles[name];style.font.name="Calibri";style.font.color.rgb=RGBColor(0,0,0)
    # Remove inherited template rules; document titles use type and whitespace.
    for style in d.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal=d.styles["Normal"];normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
    normal.paragraph_format.widow_control=True
    d.styles["Title"].font.size=Pt(25);d.styles["Title"].font.bold=True
    d.styles["Title"].paragraph_format.space_after=Pt(12)
    d.styles["Heading 1"].font.size=Pt(16);d.styles["Heading 1"].paragraph_format.space_before=Pt(8)
    d.styles["Heading 1"].paragraph_format.space_after=Pt(10)
    d.styles["Heading 2"].font.size=Pt(12)
    header=s.header.paragraphs[0];header.text="ICT4403  |  Foundations of Generative AI  |  Assignment 2"
    header.runs[0].font.size=Pt(8.5);header.runs[0].font.color.rgb=RGBColor(0,0,0)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ").font.size=Pt(8.5)
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)
    d.core_properties.title="Autoregressive next word generation"
    d.core_properties.subject="Foundations of Generative AI ICT4403 Assignment 2"
    d.core_properties.author=a.name

    d.add_paragraph("Autoregressive next word generation",style="Title")
    d.add_paragraph("Foundations of Generative AI ICT4403\nAssignment 2",style="Subtitle")
    for label,value in [("Student name",a.name),("Registration number",a.registration),
                        ("Assignment title","Foundations of Generative AI ICT4403 Assignment 2"),
                        ("System title","Alice Next Word Generator using Simple RNN BiRNN LSTM and GRU"),
                        ("Dataset","Alice's Adventures in Wonderland by Lewis Carroll"),
                        ("Submission date","16 October 2026")]:
        pp=d.add_paragraph();pp.add_run(label+": ").bold=True;pp.add_run(value)
    pp=d.add_paragraph();pp.add_run("Dataset source: ").bold=True;link(pp,meta["source"],meta["source"])
    pp=d.add_paragraph();pp.add_run("GitHub repository: ").bold=True;link(pp,a.repository,a.repository)
    heading(d,"Purpose and main finding")
    paragraph(d,"This system predicts one next word from a 16-word observed window and generates text sequentially. Simple RNN, Bidirectional RNN, LSTM and GRU were trained under one fixed protocol with three training seeds each, giving 12 completed experiments. LSTM achieved the lowest validation loss and the best held-out prediction metrics. Free-running continuations remain grammatically weak and often repetitive, so better next-word prediction does not establish fluent story generation.")
    heading(d,"Dataset overview")
    paragraph(d,"The English narrative has 12 chapters, dialogue, recurring characters and unusual vocabulary. The archived Project Gutenberg eBook 11 is accessible without an account; its source lists the work as public domain in the USA and retains the complete licensing notice. Word preprocessing yields 26,615 tokens and a training-only vocabulary of 1,167 entries including <unk>.")
    table(d,["Split","Chapters","Word tokens","Examples"],
          [[name.capitalize(),"1 to 8" if name=="train" else ("9 to 10" if name=="validation" else "11 to 12"),f"{x['word_tokens']:,}",f"{x['examples']:,}"] for name,x in meta["splits"].items()],
          [1.45,1.4,1.75,1.75])

    d.add_page_break();heading(d,"Data construction and model design")
    paragraph(d,"The pipeline removes the Gutenberg header/footer, contents and chapter headings; applies Unicode NFKC and lowercase normalization; normalizes apostrophes; and extracts ASCII words with optional internal apostrophes. Punctuation and digits are omitted and hyphenated words split. Two later-edition verse additions and their short editorial labels inside chapter 10 remain part of the archived body text.")
    paragraph(d,"Whole chapters are assigned to train, validation and test before constructing windows. For each chapter and target position t, x = words[t-16:t] and y = words[t]. Stride is one, no window crosses a chapter, and the target is outside the input. The first 16 chapter words serve as context only. Overlap occurs only within a split. A training-only frequency threshold of two defines the vocabulary; all other words map to <unk>.")
    table(d,["Model","Recurrent state","Representation","Parameters"],
          [[LABELS[r["model"]],"48 each direction" if r["model"]=="birnn" else "96 units", "96 dimensions",f"{int(r['parameters']):,}"] for r in summaries],
          [1.65,1.65,1.6,1.45])
    paragraph(d,"All architectures use a learned 64-dimensional embedding, one recurrent layer, dropout 0.15 on the context representation and a linear vocabulary classifier. The Simple RNN and BiRNN use tanh; LSTM maintains a cell state and gates; GRU uses gated state updates. Embedding and output dimensions match, while recurrent parameter counts differ. This is an equal-protocol comparison rather than an equal-parameter or equal-compute ablation.")
    heading(d,"Why bidirectional generation is valid here")
    paragraph(d,"For predicting word t, both directions of BiRNN encode only the 16 observed words ending at t-1. The reverse pass traverses this same past window backwards; it never sees the target or future text. The two final hidden states are concatenated. After sampling a word, the next past window is recomputed with that word appended. Thus the model represents p(word t | observed past window), with a bidirectional encoder of the past. A tokenwise bidirectional model with future words around its target would violate the autoregressive task; that construction is not used.")
    heading(d,"Training and checkpoint selection")
    paragraph(d,"Every model receives 18 epochs, batch size 128, AdamW learning rate 0.002, weight decay 0.0001 and gradient clipping at 1.0. Training seeds are 17, 29 and 43. A separate generator provides identical epoch permutations for all architectures within a seed. Execution uses deterministic PyTorch CPU algorithms and four threads. The lowest validation-loss epoch is saved in each run; the preferred architecture is selected by mean validation loss. Test metrics never guide tuning.")

    d.add_page_break();heading(d,"Quantitative evaluation")
    paragraph(d,"Each selected checkpoint evaluates the same 4,002 test targets from chapters 11 and 12. Loss is the token-weighted mean negative log likelihood in nats; perplexity is exp(loss); accuracy is top-1 correctness including mapped unknown words. Entries show mean and sample SD across three training seeds, not confidence intervals. Mean perplexity averages per-run perplexities.")
    def val(r,k,scale=1,digits=2):
        return f"{float(r[k+'_mean'])*scale:.{digits}f} ± {float(r[k+'_sd'])*scale:.{digits}f}"
    table(d,["Model","Val loss","Test loss","Perplexity","Accuracy %","Train s"],
          [["BiRNN" if r["model"]=="birnn" else LABELS[r["model"]],val(r,"validation_loss",digits=3),val(r,"test_loss",digits=3),val(r,"test_perplexity"),val(r,"test_accuracy",100),val(r,"training_seconds")] for r in summaries],
          [1.05,1.12,1.12,1.12,1.12,1.12])
    figure(d,root/"plots/model_comparison.png","Test perplexity, test accuracy and training time across four models with seed SD error bars")
    caption(d,"Figure 1. Mean and sample SD across three training seeds. Training seconds exclude validation, checkpoint writing and final evaluation.")
    paragraph(d,f"The add-one-smoothed training unigram baseline has test loss {baseline['loss']:.3f}, perplexity {baseline['perplexity']:.2f} and accuracy {baseline['accuracy']*100:.2f}%. All recurrent models improve on this reference. LSTM has the lowest mean validation loss (4.488), lowest test perplexity (89.67) and highest test accuracy (19.07%). GRU ranks second by loss and perplexity. Simple RNN is fastest, while BiRNN provides no observed predictive advantage under this protocol.")
    paragraph(d,"The unknown-target rate is 13.07% in test and 16.34% in validation, compared with 5.31% in training. Mean known-target test accuracy is 13.17% for RNN, 11.76% for BiRNN, 14.07% for LSTM and 13.20% for GRU. Correctly predicting <unk> contributes to overall accuracy. Perplexity concerns this mapped word space; it cannot be interpreted as an unrestricted-English perplexity.")
    paragraph(d,"Training wall time sums the epoch loops. Separate validation, evaluation and total run times are saved. All runs execute sequentially on one host in fixed architecture order, so timing includes host conditions and is not a randomized benchmark. The complete training and sample-generation experiment took 323.20 seconds; dependency installation and plotting are excluded.")

    d.add_page_break();heading(d,"Learning behavior and generalization")
    figure(d,root/"plots/learning_curves.png","Training and validation cross entropy curves across 18 epochs for four models with variability across three seeds")
    caption(d,"Figure 2. Online training loss and validation loss. Shading is sample SD over training seeds. Online training uses dropout and changing weights; validation uses evaluation mode.")
    paragraph(d,"Training loss continues falling after validation performance stops improving, indicating overfitting in the later fixed-budget epochs. The saved best epochs are 5 to 7 for Simple RNN, 4 to 5 for BiRNN, 8 to 9 for LSTM and 4 to 5 for GRU. LSTM sustains improvement longer on this corpus. Keeping the validation-selected checkpoint prevents the final overfit epoch from determining the reported held-out performance.")
    paragraph(d,"All selected-checkpoint train, validation and test metrics are recomputed in evaluation mode with dropout disabled. Source positions are disjoint across splits, but natural repeated phrases may occur in different chapters. Chapter separation and the larger held-out unknown rate make this a test of later-chapter generalization within one book, rather than a random-window interpolation task.")
    paragraph(d,"Seed SD measures training variability only. Three seeds and two held-out chapters do not support claims of universal model superiority or independent-token statistical significance. The overlapping within-split windows also make token observations dependent. Full unrounded results, selected epochs and prediction arrays are available in the repository.")

    d.add_page_break();heading(d,"Sampling temperature and text quality")
    paragraph(d,"Generation uses the validation-selected training-seed-17 checkpoint for each model, two identical 16-word training prompts, temperatures 0.5, 1.0 and 1.5, and sampling seeds 101, 202 and 303. Each continuation contains 60 words, producing 72 unedited outputs. The generator samples softmax(logits / T) over the complete vocabulary, including <unk>, and recomputes the past window after every word. No filtering or repetition penalty is applied.")
    figure(d,root/"plots/temperature_effects.png","Distinct word fraction, repeated bigram fraction and unknown word fraction versus temperature across models")
    caption(d,"Figure 3. Mean and SD over two prompts and three sampling seeds, conditional on one trained checkpoint. Surface diversity is not semantic coherence.")
    lstm_t=[r for r in temperatures if r["model"]=="lstm"]
    table(d,["Temperature","LSTM distinct 1","Repeated bigrams","Unknown words"],
          [[r["temperature"],f"{float(r['distinct_1_mean']):.3f}",f"{float(r['repeated_bigram_fraction_mean'])*100:.2f}%",f"{float(r['unk_fraction_mean'])*100:.2f}%"] for r in lstm_t],
          [1.45,1.65,1.65,1.6])
    paragraph(d,"At T=0.5, familiar local fragments appear, but sharper probabilities amplify common-word and unknown-token loops. At T=1.0, repetition drops and word variety improves; LSTM repeated bigrams fall from 29.66% to 2.54%. GRU has the strongest surface balance at this temperature, with distinct-1 0.814 and only 0.28% repeated bigrams. At T=1.5, every model has zero repeated bigrams in these samples, but the inspected sequences become loosely connected. More variety does not establish improved coherence.")
    paragraph(d,"Qualitative inspection uses the predetermined sampling-seed-101 subset for both prompts, all architectures and all temperatures (24 continuations); surface metrics use all 72. None maintains grammatical or narrative consistency for 60 words. Familiar characters recur, but scenes and references jump. LSTM's two prompts produce the identical T=1.0 continuation for RNG 101, showing rapid prompt-context loss on that trajectory. The full sample archive permits independent inspection.")
    paragraph(d,"LSTM is the preferred predictive base because of its validation and held-out metrics. T=1.0 is a reasonable exploratory compromise among the tested settings. GRU is an alternative when surface variety and modestly lower training time are priorities. No model demonstrates reliable coherent storytelling in this experiment.")

    d.add_page_break();heading(d,"Unedited generated text examples")
    paragraph(d,"The following are all four models for prompt 1, T=1.0 and sampling seed 101, with training seed 17. Each continuation is reproduced in full without correction. The common prompt is:")
    paragraph(d,meta["prompts"][0])
    for kind in LABELS:
        d.add_heading(LABELS[kind],level=2)
        sample=next(x for x in samples if x["model"]==kind and x["prompt_id"]==1 and x["temperature"]==1.0 and x["sampling_seed"]==101)
        paragraph(d,sample["continuation"])
    heading(d,"Observed failure patterns")
    paragraph(d,"The examples contain recognizable story words but broken constructions such as 'as she had asked rules', inconsistent entities and little stable narrative. At lower temperature, repeated <unk> tokens remove lexical content; at higher temperature, more rare words appear without a reliable event sequence. Similar words across models can also result from the shared categorical RNG schedule over a fixed vocabulary when probabilities are diffuse. A diversity score alone cannot distinguish useful novelty from random word choice.")
    paragraph(d,"The repository stores every model, prompt, temperature and RNG combination in Markdown, CSV and JSON, with per-sample distinct-1, distinct-2, word entropy, repeated-bigram fraction and unknown fraction. The accompanying qualitative analysis compares both prompts and the lower/higher temperature cases.")

    d.add_page_break();heading(d,"Reproducibility and limitations")
    paragraph(d,"Python source files implement data preparation, the four models, training, evaluation, generation, analysis and evidence verification. Direct dependencies are pinned in requirements.txt; the recorded complete environment, configuration and source SHA-256 hashes are saved. The raw dataset is archived with its original license and checksum, preventing an upstream edition change from silently changing results. The runner refuses to overwrite an existing run manifest.")
    paragraph(d,"Eight unit tests pass, including chapter-disjoint source positions, train-only vocabulary, target exclusion, BiRNN final-state construction, token-weighted perplexity and deterministic sampling. Verification reloads all 12 checkpoints, exactly matches their saved test prediction arrays, recomputes test metrics, confirms all 216 epoch records and selected epochs, and checks the 72 sample lengths. One saved continuation per architecture is regenerated exactly. These checks establish consistency of the delivered evidence on the recorded environment.")
    paragraph(d,"To reproduce: install requirements.txt in a Python environment; run 'python -m unittest discover -s tests -v' and 'python -m src.verify_results'. To retrain without replacing evidence, run 'python -m src.run --output results_rerun', then analyze and verify that directory. Standalone generation is available through 'python -m src.generate'. The README contains complete commands and checkpoint options. Exact numerical identity across hardware or library builds is not guaranteed.")
    paragraph(d,"Limitations include one small book, removed punctuation, only 16 words of context, lexical shift between chapters, loss of rare-word information, a small editorial component in validation, differing recurrent parameter budgets and shared hyperparameters. Seed SD is not uncertainty across books or chapters. Larger multi-book data, subword units, longer contexts and parameter-matched comparisons are future work. The implementation and writing were prepared with AI assistance; actual execution artifacts support the experimental claims.")
    heading(d,"References")
    refs=[("1. Lewis Carroll. Alice's Adventures in Wonderland. Original work 1865; Project Gutenberg eBook 11, accessed 16 October 2026.",meta["source"]),
          ("2. PyTorch 2.9 RNN documentation.","https://docs.pytorch.org/docs/2.9/generated/torch.nn.RNN.html"),
          ("3. PyTorch 2.9 LSTM documentation.","https://docs.pytorch.org/docs/2.9/generated/torch.nn.LSTM.html"),
          ("4. PyTorch 2.9 GRU documentation.","https://docs.pytorch.org/docs/2.9/generated/torch.nn.GRU.html"),
          ("5. PyTorch 2.9 cross entropy documentation.","https://docs.pytorch.org/docs/2.9/generated/torch.nn.CrossEntropyLoss.html")]
    for text,url in refs:
        pp=d.add_paragraph(text+" ");link(pp,"Source",url)
    heading(d,"Submission artifacts")
    paragraph(d,"The public repository provides implementation, supporting files, actual results, plots, generated samples and the required README. This document records the student details, assignment/system titles, dataset source and GitHub link. Its PDF export is the file to submit to the specified LMS assignment section.")
    filename=re.sub(r"[^A-Za-z0-9]+","_",a.name).strip("_")+"_"+a.registration+"_FGAI_Assignment2.docx"
    path=out/filename;d.save(path)
    shutil.copyfile(path,out/"YourName_RegNum_FGAI_Assignment2.docx")
    print(path)


if __name__=="__main__":
    main()
